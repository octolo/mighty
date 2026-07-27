"""Fill an xlsx template with the report summary, one workbook per group."""

import re
import zipfile
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

TAG = re.compile(r'{{\s*(.+?)\s*}}')


def template_content(template):
    """Accept the template as a path or as the raw bytes of the workbook."""
    if isinstance(template, (str, Path)):
        return Path(template).read_bytes()
    return template


def fill_cell(text, values, unknown):
    tag = TAG.fullmatch(text.strip())
    if tag:
        # A cell holding nothing but a tag keeps the raw value, so that Excel
        # sees a number and formulas over it still work.
        if tag.group(1) not in values:
            unknown.add(tag.group(1))
            return ''
        return values[tag.group(1)]

    def replace(match):
        if match.group(1) not in values:
            unknown.add(match.group(1))
            return ''
        return str(values[match.group(1)])

    return TAG.sub(replace, text)


def render_workbook(content, values, destination):
    """Write the template to ``destination`` with its tags replaced."""
    unknown = set()
    workbook = load_workbook(BytesIO(content))
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and '{{' in cell.value:
                    cell.value = fill_cell(cell.value, values, unknown)
    workbook.save(destination)
    return unknown


def safe_filename(name):
    return re.sub(r'[^\w.-]+', '_', name).strip('_') or 'unknown'


def render_company_workbooks(template, values_per_group, output_dir):
    """
    Write one workbook per group from ``template``.

    :param values_per_group: for each group, the value of every recap column
    :return: the written paths, and the tags the summary could not fill
    """
    content = template_content(template)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    unknown = set()
    for group, values in sorted(values_per_group.items()):
        destination = output_dir / f'{safe_filename(group)}.xlsx'
        unknown |= render_workbook(content, values, destination)
        paths.append(destination)
    return paths, unknown


def zip_workbooks(paths, destination):
    """Bundle the workbooks, since a report can cover hundreds of companies."""
    with zipfile.ZipFile(destination, 'w', zipfile.ZIP_DEFLATED) as archive:
        for path in paths:
            archive.write(path, Path(path).name)
    return destination
